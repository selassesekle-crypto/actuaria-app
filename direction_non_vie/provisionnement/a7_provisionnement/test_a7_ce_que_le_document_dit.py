# -*- coding: utf-8 -*-
"""
=============================================================================
 A7 — CE QUE LE DOCUMENT DIT, ET CE QUE LE CODE A FAIT
=============================================================================

 Ce fichier tient les lots PA-4, 5 et 6. Ils n'ont qu'une cause commune : une
 grandeur est CALCULÉE correctement et le document en dit autre chose — ou
 n'en dit rien.

 ⚠️ CE QUI EST DÉJÀ TENU AILLEURS, ET QUE CE FICHIER NE REFAIT PAS.
 `test_a7_ecretage.py` verrouille que `chain_ladder` CONSERVE la valeur brute
 de la queue et qu'il JUGE sur la brute. C'est le côté donnée, il est bon. Ce
 qui manquait est le côté DOCUMENT : `tail_brut` et `tail_ecrete` étaient
 écrits, et le relevé de tous leurs lecteurs de production en donnait ZÉRO.

 LES CINQ FAITS TENUS ICI :

 · PA-4a — un écrêtage de queue ATTEINT le commentaire, le HTML et les
   recommandations de N4. Mesuré avant : brut 1,7162 ramené à 1,5000, soit
   1 377 565 € (−18,8 %) retirés de la réserve, et zéro occurrence de la
   valeur brute dans les trois formats.
 · PA-4b — la contribution de la queue se calcule sur l'ULTIME. L'écriture
   (tail−1) × réserve_CL sous-estimait de 61,4 % sur un triangle mature.
 · Lot 5  — le retour dégradé porte les MÊMES clés que le nominal, avec les
   MÊMES types. Mesuré avant : 33 contre 16.
 · Lot 6a — le back-testing compare deux estimations du même ultime. Mesuré
   avant, sur un triangle où le provisionnement passé est exact par
   construction : ROUGE, écarts monotones +2,5 → +68,8 % avec la récence.
 · Lot 6b — le point Bootstrap est centré sur le Best Estimate, queue comprise.
   Mesuré avant : −12,7 % sous le BE, et le document imputait l'écart à « une
   possible non-normalité ».

 ⚠️ CHAQUE CLASSE PORTE SA CONTRE-ÉPREUVE : un test qui ne passe que dans un
 sens ne prouve pas que le contrôle discrimine.
=============================================================================
"""

import io
import unittest

import numpy as np

from direction_non_vie.provisionnement.a7_provisionnement.agent import (
    AgentA7Provisionnement)
from direction_non_vie.provisionnement.a7_provisionnement.n3.chain_ladder import (
    calculer_facteurs, calculer_tail_factor, projeter_ultimates)
from direction_non_vie.provisionnement.a7_provisionnement.test_a7_ibrahim import (
    RAA)


def _tri_geometrique(ratio, n=8, base=100000.0):
    """Triangle cumulé dont la cadence décroît d'un facteur `ratio` : plus
    `ratio` est proche de 1, plus la queue extrapolée est lourde."""
    m = n
    cum = np.cumsum([ratio ** j for j in range(m)])
    C = np.zeros((n, m))
    for i in range(n):
        for j in range(m):
            if i + j < n:
                C[i, j] = base * (1 + 0.03 * i) * cum[j]
    return C


def _run(C, **kw):
    d = dict(source=np.asarray(C, dtype=float), mode_declare='cumule',
             generer_graphiques=False, generer_word=False,
             n_sim_bootstrap=30, seed=42)
    d.update(kw)
    return AgentA7Provisionnement(verbose=False).run(**d)


# =============================================================================
#  PA-4a — L'ÉCRÊTAGE DE LA QUEUE ATTEINT LES LIVRABLES
# =============================================================================

class T1_La_Queue_Ecretee_Est_Dite(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.ecrete = _run(_tri_geometrique(0.97))     # tail brut > plafond
        cls.franc  = _run(_tri_geometrique(0.80))     # tail brut < plafond

    def test_un_ecretage_atteint_le_commentaire_et_le_html(self):
        tf = self.ecrete['n3']['chain_ladder']['tail_factor']
        self.assertTrue(tf.get('tail_ecrete'),
                        'le triangle de contrôle ne déclenche plus l\'écrêtage')
        brut = '%.4f' % float(tf['tail_brut'])
        for nom, txt in (('commentaire', self.ecrete.get('commentaire') or ''),
                         ('HTML', self.ecrete.get('html') or '')):
            self.assertIn(
                'ÉCRÊTÉ', txt,
                '%s : la queue est écrêtée et le document ne le dit pas — '
                'la réserve publiée est INFÉRIEURE à ce que la queue ajustée '
                'implique.' % nom)
            self.assertIn(
                brut, txt,
                '%s : la valeur brute %s n\'est pas publiée.' % (nom, brut))
        self.assertTrue(
            [r for r in (self.ecrete['n4'].get('recommandations') or [])
             if 'QUEUE ÉCRÊTÉE' in str(r)],
            'N4 n\'émet aucune recommandation sur la queue écrêtée')
        print('    OK L4-1 queue écrêtée %s → %.4f : dite au commentaire, au '
              'HTML et en recommandation N4'
              % (brut, float(tf['tail_factor'])))

    def test_sans_ecretage_le_document_n_invente_rien(self):
        """⚠️ LA CONTRE-ÉPREUVE. Un avertissement qui sort toujours ne
        distingue plus rien."""
        tf = self.franc['n3']['chain_ladder']['tail_factor']
        self.assertFalse(tf.get('tail_ecrete'),
                         'le triangle témoin est écrêté : il ne témoigne plus')
        for nom, txt in (('commentaire', self.franc.get('commentaire') or ''),
                         ('HTML', self.franc.get('html') or '')):
            self.assertNotIn(
                'ÉCRÊTÉ', txt,
                '%s : avertissement d\'écrêtage sur une queue NON écrêtée' % nom)
        print('    OK L4-2 queue non écrêtée (%.4f) : aucun avertissement'
              % float(tf['tail_factor']))


# =============================================================================
#  PA-4b — LA CONTRIBUTION DE LA QUEUE SE CALCULE SUR L'ULTIME
# =============================================================================

class T2_La_Contribution_De_La_Queue(unittest.TestCase):

    def test_la_contribution_publiee_est_celle_que_la_queue_ajoute(self):
        C = _tri_geometrique(0.90)
        f, _ = calculer_facteurs(C, 'standard')
        t = float(calculer_tail_factor(f)['tail_factor'])
        self.assertGreater(t, 1.005, 'ce triangle n\'a pas de queue à mesurer')
        r1 = float(np.sum(projeter_ultimates(C, f, tail_factor=1.0)['ibnr_brut']))
        rt = float(np.sum(projeter_ultimates(C, f, tail_factor=t)['ibnr_brut']))
        vraie = rt - r1

        r = _run(C)
        cl = r['n3']['chain_ladder']
        u = float(sum(cl.get('ultimates') or []))
        publiee = (1.0 - 1.0 / t) * u
        self.assertAlmostEqual(
            publiee / vraie, 1.0, places=2,
            msg='la contribution publiée (%.0f) n\'est pas celle que la queue '
                'ajoute (%.0f)' % (publiee, vraie))

        ancienne = (t - 1.0) * float(cl.get('reserve_totale') or 0)
        self.assertLess(
            ancienne, vraie * 0.9,
            'l\'ancienne formule (tail−1) × réserve_CL ne sous-estimait pas : '
            'la contre-épreuve est morte')
        print('    OK L4-3 contribution de la queue : %.0f € (ancienne formule '
              '%.0f €, soit %+.1f %%)'
              % (vraie, ancienne, ancienne / vraie * 100 - 100))


# =============================================================================
#  LOT 5 — LE CONTRAT DE SORTIE, ET LE FAUX ZÉRO
# =============================================================================

class T3_Le_Contrat_De_Sortie(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.nom = _run(RAA)
        cls.ech = AgentA7Provisionnement(verbose=False).run(
            source='FICHIER_QUI_N_EXISTE_PAS.csv', mode_declare='cumule',
            generer_graphiques=False, generer_word=False,
            n_sim_bootstrap=30, seed=42)

    def test_le_degrade_porte_les_memes_cles_que_le_nominal(self):
        kn, ke = set(self.nom), set(self.ech)
        self.assertEqual(
            kn - ke, set(),
            'clés ABSENTES du retour dégradé : un consommateur doit deviner '
            'si l\'absence vaut échec ou oubli')
        self.assertEqual(
            ke - kn, set(),
            'clés que le dégradé publie et que le nominal n\'a pas')
        print('    OK L5-1 contrat de sortie : %d clés des deux côtés' % len(kn))

    def test_chaque_vide_porte_le_type_du_nominal(self):
        """⚠️ « DÉCLARER LE VIDE » N'EST PAS « DÉCLARER N'IMPORTE QUEL VIDE ».
        `triangle` est une LISTE au nominal : un consommateur qui fait `len()`
        ou `np.asarray()` ne doit pas changer de branche selon le succès."""
        for cle in ('triangle', 'n1', 'n2', 'n3', 'n4', 'livrables_erreurs'):
            self.assertIs(
                type(self.ech[cle]), type(self.nom[cle]),
                '%s : %s en échec, %s au nominal'
                % (cle, type(self.ech[cle]).__name__,
                   type(self.nom[cle]).__name__))
        print('    OK L5-2 les vides portent le type du nominal')

    def test_livrables_erreurs_ne_nomme_aucun_livrable_fantome(self):
        import direction_non_vie.provisionnement.a7_provisionnement.n5_rapport as R
        cles = set(self.ech.get('livrables_erreurs') or {})
        self.assertNotIn('pdf', cles,
                         '`pdf` est un fantôme : `export_pdf` a été retiré')
        self.assertFalse(hasattr(R, 'export_pdf'),
                         'export_pdf est revenu : le plant est mort')
        self.assertIn('html', cles, '`html` EST produit et manquait')
        print('    OK L5-3 livrables_erreurs : %s' % sorted(cles))


class T4_Le_Faux_Zero_Et_Les_Etiquettes(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.r = _run(RAA, arrete='T2 2026')       # SANS exposition : BF et CC
        cls.wb = None                             # exclues, donc non calculées

    def _classeur(self):
        from openpyxl import load_workbook
        if self.wb is None:
            type(self).wb = load_workbook(io.BytesIO(self.r['excel_bytes']))
        return self.wb

    def test_aucun_onglet_ne_publie_zero_pour_une_methode_non_calculee(self):
        wb = self._classeur()
        for feuille in (wb.sheetnames[0], wb.sheetnames[3]):
            for row in wb[feuille].iter_rows(values_only=True):
                if not row or not any(isinstance(c, str) and 'Bornhuetter' in c
                                      for c in row if c is not None):
                    continue
                vals = [c for c in row if c is not None]
                self.assertIn(
                    'non calculée', [str(v) for v in vals],
                    '%s publie %r pour une méthode non calculée' % (feuille, vals))
                print('    OK L5-4 [%s] %s' % (feuille, vals))

    def test_le_classeur_distingue_l_arrete_de_la_date_de_generation(self):
        ws = self._classeur().worksheets[0]
        etiquettes = {str(row[0]) for row in ws.iter_rows(values_only=True)
                      if row and isinstance(row[0], str)}
        self.assertIn('Arrêté', etiquettes,
                      'le classeur n\'a AUCUN champ d\'arrêté')
        self.assertIn('Généré le', etiquettes,
                      'la date de génération a disparu avec l\'étiquette fautive')
        self.assertNotIn('Date rapport', etiquettes,
                         '« Date rapport » portait deux notions sous un seul nom')
        print('    OK L5-5 le classeur porte « Arrêté » ET « Généré le »')

    def test_le_compte_d_onglets_ne_vit_plus_dans_la_prose(self):
        import direction_non_vie.provisionnement.a7_provisionnement.n5_excel as X
        src = io.open(X.__file__, encoding='utf-8').read()
        self.assertNotIn(
            '10 onglets', src,
            'le compte est recopié dans la prose : Munich CL étant '
            'conditionnel, AUCUN nombre fixe n\'est vrai')
        self.assertGreaterEqual(len(self._classeur().worksheets), 10)
        print('    OK L5-6 %d feuilles produites, aucun compte en prose'
              % len(self._classeur().worksheets))

    def test_bfcc_h6_figure_au_jugement_publie(self):
        jug = str((self.r['n4'] or {}).get('jugement') or '')
        manquants = [c for c in ('BFCC-H1', 'BFCC-H2', 'BFCC-H3', 'BFCC-H4',
                                 'BFCC-H5', 'BFCC-H6') if c not in jug]
        self.assertEqual(
            manquants, [],
            'hypothèses absentes du jugement publié : %s. BFCC-H6 est '
            'BLOQUANTE pour Cape Cod (`_HYPOTHESES_BLOQUANTES`).' % manquants)
        print('    OK L5-7 les six BFCC figurent au jugement publié')


# =============================================================================
#  LOT 6 — LA RÉFÉRENCE DE COMPARAISON
# =============================================================================

class T5_Le_Back_Testing_Compare_Deux_Ultimes(unittest.TestCase):

    @staticmethod
    def _triangle_exact(n=8):
        """Motif chain-ladder STRICTEMENT multiplicatif : Chain Ladder y est
        exact par construction, donc le vrai boni/mali de liquidation est nul.
        Tout écart publié sur ce triangle est un artefact du module."""
        fv = [1.60, 1.28, 1.14, 1.08, 1.045, 1.025, 1.012]
        cad = [1.0]
        for x in fv:
            cad.append(cad[-1] * x)
        cad = np.array(cad) / cad[-1]
        C = np.zeros((n, n))
        for i in range(n):
            for j in range(n):
                if i + j < n:
                    C[i, j] = 10000.0 * (1 + 0.02 * i) * cad[j]
        return C

    def test_un_provisionnement_exact_ne_publie_pas_un_ecart_qui_croit(self):
        from direction_non_vie.provisionnement.a7_provisionnement.n3.backtesting \
            import calculer_backtesting
        bt = calculer_backtesting(self._triangle_exact())
        ecarts = [r['ecart_pct_n1'] for r in (bt.get('tableau') or [])
                  if r.get('ecart_pct_n1') is not None]
        self.assertGreaterEqual(len(ecarts), 4, 'trop peu d\'années évaluées')
        etendue = max(ecarts) - min(ecarts)
        self.assertLess(
            etendue, 2.0,
            'les écarts s\'étalent sur %.1f points (%s) : c\'est la signature '
            'du DÉVELOPPEMENT RESTANT, qui croît avec la récence, pas celle '
            'd\'un écart de provisionnement.' % (etendue, ecarts))
        self.assertNotEqual(
            bt.get('statut'), 'ROUGE',
            'ROUGE publié sur un triangle où Chain Ladder est exact par '
            'construction')
        print('    OK L6-1 triangle exact : statut %s, écarts %s (étendue '
              '%.1f pt)' % (bt.get('statut'), ecarts, etendue))

    def test_le_total_dit_la_meme_chose_que_ses_lignes(self):
        """⚠️ LE TOTAL SE CALCULE SUR LA POPULATION COMMUNE. Sommer les
        ultimes N−1 des années qui en ont un et les comparer à une base sommée
        sur TOUTES les années fabrique un écart qui n'est celui d'aucune
        ligne."""
        from direction_non_vie.provisionnement.a7_provisionnement.n3.backtesting \
            import calculer_backtesting
        bt = calculer_backtesting(self._triangle_exact())
        lignes = [r['ecart_pct_n1'] for r in (bt.get('tableau') or [])
                  if r.get('ecart_pct_n1') is not None]
        total = (bt.get('totaux') or {}).get('ecart_pct_n1')
        self.assertIsNotNone(
            total, 'la ligne des totaux ne porte plus d\'écart : `ultimate_n0` '
                   'est-il écrit dans `tableau` et non dans `alertes` ?')
        self.assertLessEqual(
            abs(total - sum(lignes) / len(lignes)), 1.0,
            'le total annonce %.1f %% quand ses lignes disent %.1f %% en '
            'moyenne' % (total, sum(lignes) / len(lignes)))
        print('    OK L6-2 total %.1f %% ≡ moyenne des lignes %.1f %%'
              % (total, sum(lignes) / len(lignes)))


class T6_Le_Bootstrap_Est_Centre(unittest.TestCase):

    def test_le_point_bootstrap_est_centre_sur_le_be_queue_comprise(self):
        rng = np.random.default_rng(4)
        n = m = 8
        cum = np.cumsum([0.90 ** j for j in range(m)])
        C = np.zeros((n, m))
        for i in range(n):
            for j in range(m):
                if i + j < n:
                    C[i, j] = 100000.0 * (1 + 0.03 * i) * cum[j] * \
                        float(rng.uniform(0.96, 1.05))
                if j > 0 and i + j < n:
                    C[i, j] = max(C[i, j], C[i, j - 1] * 1.001)
        r = _run(C, n_sim_bootstrap=400)
        tail = float(r['n3']['chain_ladder']['tail_factor']['tail_factor'])
        self.assertGreater(tail, 1.005, 'ce triangle n\'a pas de queue')
        be = float(r['n4']['best_estimate'])
        pt = float((r['n3'].get('bootstrap') or {}).get('be_bootstrap') or 0)
        self.assertLess(
            abs(pt / be - 1.0), 0.02,
            'point Bootstrap %.0f contre BE %.0f (%.1f %%) : la cible de '
            'recentrage ne porte pas la queue que le BE porte.'
            % (pt, be, (pt / be - 1) * 100))
        print('    OK L6-3 queue %.4f : point Bootstrap %.0f ≡ BE %.0f '
              '(%+.2f %%)' % (tail, pt, be, (pt / be - 1) * 100))


if __name__ == '__main__':
    unittest.main(verbosity=2)
