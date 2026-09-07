"""LE LITTERAL NEUTRE A LA PLACE DE L'ABSENCE -- cause (a) de l'audit du
06/09/2026, quatre constats.

Une valeur que PERSONNE N'A MESUREE occupe la place d'une mesure, et elle se
lit comme une mesure. C'est la cause qui produit le plus de constats PUBLIES
du peri1metre, et aucun de ses correctifs ne deplace un prix.

  LN-1  `A6-1` -- la piste d'audit ne porte pas 0 quand le moteur a dit ne
        pas savoir, et le document signe non plus ;
  LN-2  `A6-1` -- ET LA PROPRIETE SE VERIFIE SUR TOUTE LA PISTE D'AUDIT, PAS
        SUR LA SEULE CLE CORRIGEE : l'assiette est DERIVEE de l'AST du
        dictionnaire, aucune liste de cles n'est tenue a la main ;
  LN-3  `EX-2` -- un A/E MESURE a 0.0 se publie 0.0, et une ABSENCE se publie
        comme une absence. Les deux sens, parce que `or` confond les deux ;
  LN-4  `A1-2` -- le classeur A1 prive de `qualite` ne publie ni zero ni
        pastille verte ;
  LN-5  `A4-1` -- la validation ML ne PLANTE plus quand le PSI n'a pas pu
        etre mesure ;
  LN-6  `A1-2` -- les DEUX exportateurs A1 tiennent la meme doctrine : c'est
        la propriete qui empeche la surface jumelle de diverger a nouveau.

⚠️⚠️ POURQUOI LN-2 EXISTE A COTE DE LN-1. Corriger `ae_ratio` seul refermait
l'instance et laissait la classe ouverte : rien n'aurait empeche la
prochaine cle de piste d'audit de naitre avec un `.get(cle, 0)`. LN-2 lit
l'arbre syntaxique du dictionnaire et refuse TOUT defaut litteral non nul,
sur les cles d'aujourd'hui comme sur celles de demain.

✅ CE QUI ETAIT SIGNALE ICI EST DESORMAIS CORRIGE -- constat `MES-1`, ferme le
07/09/2026, voir `test_absence_pas_verdict.py`. Cette note disait :
<< `conformite:908` rend "BIAIS DE TARIFICATION -- A/E walk-forward = None"
quand l'A/E est absent ; la BRANCHE tire, mais l'ETAT n'a pas ete produit sur
trois portefeuilles reels. On ne corrige pas un etat qu'on n'a pas su
produire. Signale, non traite. >>

⚠️⚠️ ET C'EST L'ATTENTE DE L'ETAT QUI ETAIT LE MAUVAIS CRITERE. Le meme lot
avait ajoute une quatrieme branche censee dire l'absence -- placee APRES celle
qui court-circuite, elle etait du CODE MORT, et son propre sceau ne l'a pas
vu parce qu'il assemblait le `backtest` a la main. *Exiger un portefeuille
reel avant de corriger a laisse le defaut vivre ET son correctif mourir ;
c'est la LECTURE DES BRANCHES, pas la production de l'etat, qui l'a tranche.*

Tout en `unittest.TestCase` : la gate lance `unittest discover`.
"""
import ast
import io
import os
import sys
import unittest
import warnings

_ICI = os.path.dirname(os.path.abspath(__file__))
_RACINE = os.path.dirname(os.path.dirname(_ICI))
for _c in (_RACINE, _ICI):
    if _c not in sys.path:
        sys.path.insert(0, _c)

from direction_non_vie.tarification.services import tarif_excel as TX

_A6 = os.path.join(_ICI, 'a6_comparaison', 'agent.py')


def _cellules(octets, feuille=0, jusqu_a=30):
    from openpyxl import load_workbook
    classeur = load_workbook(io.BytesIO(octets), data_only=True)
    feuille = classeur[classeur.sheetnames[feuille]]
    lues = []
    for ligne in feuille.iter_rows(min_row=1, max_row=jusqu_a):
        for cellule in ligne:
            if cellule.value not in (None, ''):
                couleur = getattr(getattr(cellule.fill, 'fgColor', None),
                                  'rgb', None)
                lues.append((cellule.coordinate, cellule.value,
                             couleur if isinstance(couleur, str) else None))
    return lues


#: Le vert de la charte, tel que `_kpi` le pose. ⚠️ Releve sur le classeur
#: REELLEMENT produit, pas recopie d'une constante : c'est cette valeur-la
#: qu'un lecteur voit.
_VERT_CHARTE = '002ECC71'

#: Un `result_a1` qui a REUSSI mais ne porte aucune qualite. C'est la forme
#: exacte que la fonction reexportee peut recevoir.
_A1_SANS_QUALITE = {'success': True, 'audit_id': 'LN-4', 'statut_rag': 'VERT'}

#: ⚠️⚠️ DEUX DEFAUTS LITTERAUX TROUVES PAR LN-2 ET RENVOYES AU LOT 3, avec
#: leur motif — ce ne sont PAS des faux positifs.
#:
#:   `stabilite_wf` et `modele_production` prennent `''` comme defaut. La
#:   difference avec `ae_ratio` est mesurée : `0` FABRIQUE UNE MESURE (un A/E
#:   de 0 se lit « le modele attend des sinistres et on n'en observe aucun »),
#:   `''` rend une case BLANCHE — ce qui ne ment pas sur une grandeur, mais ne
#:   declare pas l'absence non plus.
#:
#:   ⚠️ ET LES CORRIGER ICI SERAIT LE MAUVAIS LOT. Mesure du 06/09/2026 :
#:   leurs consommateurs lisent `backtest`, pas la piste d'audit, et chacun
#:   applique SON PROPRE defaut — `'—'` (`rapport_modeles:1515` et `:2204`),
#:   `'N/A'` (`tarif_excel:738`, `a6:2964`), `''` (`rapport_modeles:2726`,
#:   `conformite:903`). Choisir un rendu unique entre quatre surfaces est la
#:   CAUSE (b), donc le lot 3.
_EXEMPTES_LOT_3 = {'stabilite_wf', 'modele_production'}


# =============================================================================
#  LN-1 — LE COMPORTEMENT : UN SEUL EXERCICE, ET LA PISTE D'AUDIT SE TAIT
# =============================================================================

class TestPisteAuditUnSeulExercice(unittest.TestCase):
    """⚠️⚠️ LE DECLENCHEUR EST LE CAS NORMAL D'UNE PREMIERE TARIFICATION : un
    client qui ne fournit qu'un exercice. `_backtesting_temporel` retourne
    alors un dictionnaire a deux cles, SANS `ae_ratio` -- et le `.get(..., 0)`
    convertissait cette absence en mesure, dans le chapitre qu'un controleur
    lit comme la piste d'audit.
    """

    @classmethod
    def setUpClass(cls):
        warnings.filterwarnings('ignore')
        chemin = os.path.join(_ICI, 'a6_comparaison')
        if chemin not in sys.path:
            sys.path.insert(0, chemin)
        import test_a6_comparaison as fixtures

        from direction_non_vie.tarification.a6_comparaison.agent import (
            AgentA6Comparaison,
        )
        r_a2 = fixtures._make_r_a2_avec_annee(600)
        # ⚠️ UN SEUL EXERCICE : le walk-forward ne peut pas tourner, et c'est
        # exactement le chemin ou le moteur ne pose jamais `ae_ratio`.
        r_a2['dataframe']['annee_souscription'] = 2024
        cls.resultat = AgentA6Comparaison(
            models_path='/tmp', audit_path='/tmp', verbose=False).run(
            result_a2=r_a2, result_a3=fixtures._make_r_a3(),
            result_a4=fixtures._make_r_a4(), col_cible='nb_sinistres',
            generer_graphiques=False)

    def test_LN1_le_moteur_ne_pose_pas_ae_ratio_sur_ce_chemin(self):
        """L'assiette se declare : sans ce controle, LN-1b resterait vert le
        jour ou le moteur se mettrait a poser une valeur -- il ne mesurerait
        plus le defaut, il constaterait un accord."""
        backtest = self.resultat.get('backtest') or {}
        self.assertNotIn(
            'ae_ratio', backtest,
            "le moteur pose desormais `ae_ratio` sur ce chemin : LN-1b ne "
            f'mesure plus le defaut. Cles rendues : {sorted(backtest)}')
        print(f'    LN-1 moteur, un seul exercice : cles {sorted(backtest)}')

    def test_LN1b_la_piste_d_audit_publie_une_ABSENCE_jamais_un_zero(self):
        trace = self.resultat.get('audit_trail') or {}
        self.assertIn('ae_ratio', trace)
        self.assertIsNone(
            trace['ae_ratio'],
            "la piste d audit publie un LITTERAL la ou le moteur a dit ne pas "
            f"savoir : {trace['ae_ratio']!r}. Un A/E de 0 se lit << le modele "
            "attend des sinistres et on n en observe aucun >>.")
        print('    LN-1b piste d audit : absence, pas un zero')


# =============================================================================
#  LN-2 — LA PROPRIETE, DERIVEE DE L'AST (aucune liste tenue a la main)
# =============================================================================

class TestPisteAuditSansLitteral(unittest.TestCase):

    def test_LN2_aucune_cle_de_piste_d_audit_n_a_de_defaut_litteral(self):
        """⚠️⚠️ L'ASSIETTE EST DERIVEE : on lit le dictionnaire
        `_audit_trail_a6` dans l'arbre syntaxique et on examine CHACUNE de ses
        cles. Une liste ecrite ici serait la meme dette que celle qu'on ferme
        -- elle divergerait le jour ou une cle est ajoutee.
        """
        with open(_A6, encoding='utf-8') as fichier:
            arbre = ast.parse(fichier.read())

        dictionnaire = None
        for noeud in ast.walk(arbre):
            if (isinstance(noeud, ast.Assign)
                    and any(isinstance(c, ast.Name)
                            and c.id == '_audit_trail_a6'
                            for c in noeud.targets)
                    and isinstance(noeud.value, ast.Dict)):
                dictionnaire = noeud.value
                break
        self.assertIsNotNone(
            dictionnaire,
            "`_audit_trail_a6` n'est plus un dictionnaire litteral : ce "
            "controle ne mesure plus rien et doit etre reecrit, pas retire.")

        fautifs = []
        examinees = 0
        for cle, valeur in zip(dictionnaire.keys, dictionnaire.values):
            nom = cle.value if isinstance(cle, ast.Constant) else '<calculee>'
            for noeud in ast.walk(valeur):
                if not (isinstance(noeud, ast.Call)
                        and isinstance(noeud.func, ast.Attribute)
                        and noeud.func.attr == 'get'
                        and len(noeud.args) == 2):
                    continue
                examinees += 1
                defaut = noeud.args[1]
                if isinstance(defaut, ast.Constant) and defaut.value is not None:
                    fautifs.append((nom, defaut.value, noeud.lineno))
        self.assertGreaterEqual(
            examinees, 1,
            'aucun `.get(cle, defaut)` trouve dans la piste d audit : '
            'le controle porterait sur une assiette vide')

        # ⚠️ L'EXEMPTION EST NOMMEE, ET ELLE SE RE-VERIFIE. Une categorie
        # (<< les chaines vides passent >>) aurait laisse entrer le prochain
        # `.get(cle, 'N/A')` sans un mot ; ces deux cles-la sont citees une par
        # une, avec leur motif, et le controle EXIGE qu'elles existent encore.
        restants = [f for f in fautifs if f[0] not in _EXEMPTES_LOT_3]
        self.assertEqual(
            restants, [],
            'une cle de piste d audit a un defaut LITTERAL non declare : le '
            'moteur qui dit ne pas savoir sera publie comme ayant mesure. '
            f'{restants}')
        # ⚠️ UNE EXEMPTION QUI SURVIT A SON SUJET EST UNE JUSTIFICATION SANS
        # OBJET : si l'une des deux est corrigee, la retirer d'ici.
        orphelines = _EXEMPTES_LOT_3 - {f[0] for f in fautifs}
        self.assertEqual(
            orphelines, set(),
            f'exemption(s) sans objet -- a retirer de `_EXEMPTES_LOT_3` : '
            f'{sorted(orphelines)}')
        print(f'    LN-2 piste d audit : {len(dictionnaire.keys)} cles, '
              f'{examinees} `.get` a defaut examines, 0 litteral non declare '
              f'({len(_EXEMPTES_LOT_3)} exemptes, renvoyes au lot 3)')


# =============================================================================
#  LN-3 — `or` CONFOND L'ABSENCE ET LE ZERO MESURE
# =============================================================================

class TestAERatioMesureAZero(unittest.TestCase):
    """⚠️ LES DEUX SENS. Un controle qui n'exercerait que l'absence resterait
    vert sur `or` ; un controle qui n'exercerait que le zero mesure resterait
    vert sur une garde qui masque tout.
    """

    def _html(self, ae):
        from direction_non_vie.tarification.services import (
            rapport_equipe_tarif as RE,
        )
        resultats = {
            'a6': {'success': True, 'statut_rag': 'VERT',
                   'modele_production': {'modele': 'GLM_POISSON',
                                         'score_global': 0.81},
                   'classement': [],
                   'backtest': {'disponible': True, 'ae_ratio': ae,
                                'gini_wf_moyen': 0.1783},
                   'audit_trail': {'profil_ponderation': 'equilibre'}},
        }
        octets = RE.generer_rapport_equipe_tarification(
            resultats, branche='non_vie', arrete='2026-06-30',
            audit_id='LN-3', formats=['html']).get('html_bytes') or b''
        return octets.decode('utf-8', 'replace')

    def test_LN3_un_AE_mesure_a_zero_est_publie_comme_une_mesure(self):
        """⚠️⚠️ MESURE, PAS THEORIQUE : une derniere fenetre de walk-forward
        sans aucun sinistre observe (150 contrats, 0 sinistre) rend
        `ae_ratio = 0.0`. C'est le cas normal d'un exercice recent.
        """
        texte = self._html(0.0)
        self.assertNotIn(
            'non calcule', texte,
            "un A/E MESURE a 0.0 est publie << non calcule >> : `or` teste la "
            "veracite, pas l absence")
        self.assertIn('A/E ratio', texte)
        print('    LN-3 A/E mesure a 0.0 : publie comme une mesure')

    def test_LN3b_une_absence_reste_declaree_comme_une_absence(self):
        """Controle NEGATIF declare : sans lui, une garde qui masque tout
        passerait le test precedent."""
        self.assertIn(
            'non calcule', self._html(None),
            "un A/E ABSENT doit rester declare : une garde qui publierait "
            '0 aurait ferme un defaut en en ouvrant son inverse')
        print('    LN-3b A/E absent : declare comme une absence')


# =============================================================================
#  LN-4, LN-6 — LE CLASSEUR A1, ET SA SURFACE JUMELLE
# =============================================================================

class TestClasseurA1SansQualite(unittest.TestCase):

    def test_LN4_aucun_zero_ni_pastille_verte_sans_qualite(self):
        """⚠️⚠️ SUR LES OCTETS PRODUITS, jamais sur le dictionnaire interne.
        Mesure avant correctif : sept zeros et deux << Conforme >> verts, dont
        un qui survivait meme a l'absence de `statut_rag`.
        """
        octets = TX.export_excel_a1(dict(_A1_SANS_QUALITE), 'LN-4')
        self.assertTrue(octets, 'aucun classeur produit : rien n est mesure')
        lues = _cellules(octets)
        libelles = [str(v) for _, v, _ in lues]
        self.assertTrue(
            any('Qualité du fichier' in x for x in libelles),
            f'la phrase d absence n atteint pas le classeur : {libelles[:14]}')

        zeros = [(c, v) for c, v, _ in lues
                 if isinstance(v, (int, float)) and not isinstance(v, bool)
                 and v == 0]
        self.assertEqual(
            zeros, [],
            f'le classeur publie {len(zeros)} zero(s) sur un fichier jamais '
            f'transmis : {zeros}')

        verts = [(c, v) for c, v, couleur in lues
                 if couleur == _VERT_CHARTE]
        self.assertEqual(
            verts, [],
            f'le classeur publie {len(verts)} pastille(s) VERTE(s) sur un '
            f'fichier jamais transmis : {verts}')
        print(f'    LN-4 classeur A1 sans qualite : 0 zero, 0 pastille verte, '
              f'{len(lues)} cellules lues')

    def test_LN6_les_DEUX_exportateurs_A1_tiennent_la_meme_doctrine(self):
        """⚠️⚠️ C'EST LA PROPRIETE, PAS L'INSTANCE. Le correctif de `A6.7`
        avait atteint `rapport_equipe_tarif` et pas `tarif_excel` : les deux
        surfaces publient le meme fait, une seule le declarait. Ce controle
        les compare, il ne verifie pas deux fois la meme.
        """
        from direction_non_vie.tarification.services import (
            rapport_equipe_tarif as RE,
        )
        octets_a1 = TX.export_excel_a1(dict(_A1_SANS_QUALITE), 'LN-6')
        octets_eq = RE.generer_rapport_equipe_tarification(
            {'a1': dict(_A1_SANS_QUALITE)}, branche='non_vie',
            arrete='2026-06-30', audit_id='LN-6',
            formats=['excel']).get('excel_bytes') or b''
        self.assertTrue(octets_a1 and octets_eq,
                        'un des deux classeurs est vide : rien n est compare')

        def porte_l_absence(octets):
            """⚠️ TOUTES LES FEUILLES. Ma premiere version ne lisait que la
            premiere : le bloc qualite du rapport d'equipe vit en feuille 2,
            et le controle rougissait sur un defaut de SONDE."""
            from openpyxl import load_workbook

            from core.conformite_reglementaire import NON_TRANSMIS
            classeur = load_workbook(io.BytesIO(octets), data_only=True)
            for nom in classeur.sheetnames:
                for ligne in classeur[nom].iter_rows():
                    for cellule in ligne:
                        if NON_TRANSMIS in str(cellule.value or ''):
                            return True
            return False

        for nom, octets in (('export_excel_a1', octets_a1),
                            ('rapport_equipe (Excel)', octets_eq)):
            with self.subTest(surface=nom):
                self.assertTrue(
                    porte_l_absence(octets),
                    f'{nom} ne declare pas l absence de qualite : la surface '
                    f'jumelle a diverge a nouveau')
        print('    LN-6 les deux exportateurs A1 declarent l absence')


# =============================================================================
#  LN-5 — LA VALIDATION ML NE PLANTE PLUS SUR UN PSI NON MESURE
# =============================================================================

class TestValidationMLSansPSI(unittest.TestCase):

    def test_LN5_un_PSI_non_mesure_ne_fait_plus_lever_la_validation(self):
        """⚠️⚠️ CE CONSTAT NE PUBLIE PAS UN CHIFFRE FAUX : IL EMPECHE TOUTE
        PUBLICATION. `round(None, 4)` leve `TypeError`, et selon le `try`
        englobant la validation actuarielle du modele est perdue ENTIERE.
        """
        warnings.filterwarnings('ignore')
        from core.conformite_reglementaire import NON_MESURE
        from direction_non_vie.tarification.a4_ml.agent import AgentA4ML

        agent = AgentA4ML.__new__(AgentA4ML)
        classement = [{'modele': 'gbm', 'gini_test': 0.21, 'gini_train': 0.25,
                       'rmse_test': 0.44, 'overfit_ratio': 1.19}]
        # X_train / X_test absents => `_psi_reel` ne mesure rien ;
        # `monitoring['psi']` absent aussi => le dernier recours est vide.
        resultat = agent._valider_modele_ml(
            classement=classement, monitoring={'psi': None},
            n_train=800, n_test=200)

        self.assertIn('h2_psi', resultat)
        self.assertEqual(
            resultat['h2_psi']['psi'], NON_MESURE,
            "un PSI non mesure doit se DECLARER, jamais valoir un nombre : "
            f"publie {resultat['h2_psi']['psi']!r}")
        self.assertEqual(resultat['h2_psi']['statut'], 'AMBRE')
        self.assertIn(
            NON_MESURE, resultat['h2_psi']['titre_graphique'],
            'le titre de figure fabrique encore un nombre : '
            f"{resultat['h2_psi']['titre_graphique']}")
        print('    LN-5 PSI non mesure : validation rendue, absence declaree')


if __name__ == '__main__':
    unittest.main(verbosity=2)
