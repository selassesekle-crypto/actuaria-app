"""
Le rapport actuariel COMBINÉ — l'ouvrir, et pas seulement le peser.

⚠️ POURQUOI CE FICHIER EXISTE, ET CE QUE LA MESURE A MONTRÉ.
Le périmètre compte trois producteurs de livrables signés. `test_rapports_sp`
en couvre deux — les rapports Santé et Prévoyance — en OUVRANT réellement le
`.docx`. Le troisième, `rapport_actuariel`, est celui que signe l'actuaire, et
le seul test qui le touchait se contentait de :

    assert isinstance(r["excel_bytes"], bytes)
    assert len(r["excel_bytes"]) > 1000

Un classeur dont l'onglet « Provisions » sort entièrement vide passe ces deux
assertions. C'est exactement ce qui se produisait : la ligne « TP = BE + RA »
affichait « — » dans les trois colonnes, parce que `tp_sante` et `tp_prev`
n'étaient produits par personne. Peser un livrable n'est pas le lire.

⚠️ CE QUE CE SCEAU A TROUVÉ EN S'ÉCRIVANT, et que ni l'audit ni moi n'avions vu.
En ouvrant l'onglet « Sinistralité », la colonne « Remb. SS (%) » affichait :

    Medecine 1 756,0 %   Pharmacie 1 258,0 %   Hospitalisation 246 400,0 %
    Dentaire 3 960,0 %   Optique 0,0 %

S1 publie `remb_ss` = `cout_acte x tc_ss`, donc un MONTANT EN EUROS. Le
classeur le multipliait par 100 et l'intitulait « % ». Après correction, la
colonne rend les taux réels de la Sécurité sociale : 70 %, 65 %, 80 %, 25 %,
0 % — ceux des tables DREES. Le défaut était sur un document signé, et aucune
mesure d'octets ne pouvait le voir.

CE QUE CE SCEAU VÉRIFIE
  1. le classeur s'OUVRE et porte ses huit onglets ;
  2. l'onglet « Provisions » porte des montants, la ligne TP comprise ;
  3. les taux de remboursement restent dans [0, 1] — une borne, pas une valeur ;
  4. les parts de charge somment à 100 % ;
  5. le Word s'OUVRE, et son texte inclut celui de ses TABLEAUX
     (`doc.paragraphs` les ignore — piège déjà payé une fois).
"""
import io
import re
import unittest


def _nombre(cellule):
    """Lit « 137,553 » ou « 70.0% » sans jamais lever."""
    if cellule is None:
        return None
    texte = str(cellule).strip().replace(",", "").replace(" ", "")
    pourcent = texte.endswith("%")
    texte = texte.rstrip("%").replace("€", "").strip()
    if not re.fullmatch(r"-?\d+(\.\d+)?", texte):
        return None
    valeur = float(texte)
    return valeur / 100.0 if pourcent else valeur


class TestLeClasseurCombineSOuvreEtPorteSesMontants(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        import contextlib

        from .prevoyance.p1_tarification.agent import (
            AgentP1TarificationPrevoyance)
        from .prevoyance.p2_tables_morbidite.agent import AgentP2TablesMorbidite
        from .prevoyance.p3_provisionnement.agent import (
            AgentP3ProvissionnementPrevoyance)
        from .prevoyance.p4_reporting.agent import AgentP4ReportingPrevoyance
        from .rapport_actuariel.agent import AgentSPRapportActuariel
        from .sante.s1_tarification.agent import AgentS1TarificationSante
        from .sante.s2_provisionnement.agent import AgentS2ProvissionnementSante
        from .sante.s3_reporting.agent import AgentS3ReportingSante

        with contextlib.redirect_stdout(io.StringIO()):
            r1 = AgentP1TarificationPrevoyance(verbose=False).run(
                age=40, salaire_brut=45_000, categorie="employe",
                generer_graphiques=False)
            r2 = AgentP2TablesMorbidite(verbose=False).run(
                result_p1=r1, generer_graphiques=False)
            r3 = AgentP3ProvissionnementPrevoyance(verbose=False).run(
                result_p1=r1, result_p2=r2, generer_graphiques=False)
            r4 = AgentP4ReportingPrevoyance(verbose=False).run(
                result_p1=r1, result_p2=r2, result_p3=r3,
                fonds_propres=5_000_000, generer_graphiques=False)
            s1 = AgentS1TarificationSante(verbose=False).run(
                generer_graphiques=False)
            s2 = AgentS2ProvissionnementSante(verbose=False).run(
                result_s1=s1, generer_graphiques=False)
            s3 = AgentS3ReportingSante(verbose=False).run(
                result_s1=s1, result_s2=s2, fonds_propres=5_000_000,
                generer_graphiques=False)
            cls.rapport = AgentSPRapportActuariel(verbose=False).run(
                result_s1=s1, result_s2=s2, result_s3=s3, result_p1=r1,
                result_p2=r2, result_p3=r3, result_p4=r4,
                generer_graphiques=False)

    def _classeur(self):
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(self.rapport["excel_bytes"]))

    def _lignes(self, nom_onglet):
        ws = self._classeur()[nom_onglet]
        return list(ws.iter_rows(values_only=True))

    def test_la_chaine_va_au_bout(self):
        self.assertTrue(self.rapport["success"], self.rapport.get("erreur"))

    def test_le_classeur_s_ouvre_et_porte_ses_onglets(self):
        """Peser un livrable n'est pas le lire."""
        onglets = self._classeur().sheetnames
        self.assertGreaterEqual(
            len(onglets), 8,
            "Le classeur signe doit porter ses huit onglets ; trouves : %s"
            % onglets)
        for attendu in ("3. Provisions", "2. Sinistralité"):
            with self.subTest(onglet=attendu):
                self.assertIn(attendu, onglets)

    def test_l_onglet_provisions_porte_des_montants(self):
        lignes = self._lignes("3. Provisions")
        montants = 0
        for ligne in lignes[2:]:
            for cellule in ligne[1:4]:
                valeur = _nombre(cellule)
                if valeur is not None and valeur > 0:
                    montants += 1
        self.assertGreater(
            montants, 8,
            "L'onglet Provisions ne porte que %d montant(s) : un classeur de "
            "provisions techniques vide passerait `len(bytes) > 1000`."
            % montants)

    def test_la_ligne_des_provisions_techniques_n_est_pas_vide(self):
        """C'était le défaut : « TP = BE + RA » sortait vide sur trois colonnes."""
        ligne_tp = None
        for ligne in self._lignes("3. Provisions"):
            if ligne and ligne[0] and "TP" in str(ligne[0]):
                ligne_tp = ligne
                break
        self.assertIsNotNone(ligne_tp, "La ligne « TP = BE + RA » a disparu.")

        sante, prev, consolide = (_nombre(ligne_tp[1]), _nombre(ligne_tp[2]),
                                  _nombre(ligne_tp[3]))
        for etiquette, valeur in (("Sante", sante), ("Prevoyance", prev),
                                  ("Consolide", consolide)):
            with self.subTest(colonne=etiquette):
                self.assertIsNotNone(
                    valeur,
                    "Colonne %s de la ligne TP : %r. Les provisions techniques "
                    "d'un rapport signe ne peuvent pas sortir vides."
                    % (etiquette, ligne_tp))
                self.assertGreater(valeur, 0)
        self.assertAlmostEqual(
            sante + prev, consolide, delta=max(consolide * 0.001, 1.0),
            msg="Le consolide doit etre la somme des deux branches : "
                "%s + %s != %s" % (sante, prev, consolide))

    def test_les_taux_de_remboursement_restent_des_taux(self):
        """⚠️ Le défaut trouvé ici : 246 400 % sur l'hospitalisation.

        La borne porte sur [0, 1] et non sur une valeur attendue : elle tient
        quelles que soient les tables, et c'est ce qui la rend utile.
        """
        lignes = self._lignes("2. Sinistralité")
        entetes = [str(c) if c else "" for c in lignes[1]]
        self.assertIn("Remb. SS (%)", entetes)
        colonne = entetes.index("Remb. SS (%)")

        vus = 0
        for ligne in lignes[2:]:
            taux = _nombre(ligne[colonne])
            if taux is None:
                continue
            vus += 1
            with self.subTest(poste=ligne[0]):
                self.assertGreaterEqual(taux, 0.0)
                self.assertLessEqual(
                    taux, 1.0,
                    "Taux de remboursement de %s = %.1f%% : la colonne "
                    "annonce un taux et publie un MONTANT multiplie par 100."
                    % (ligne[0], taux * 100))
        self.assertGreater(vus, 2, "Moins de trois postes lus : assiette vide.")

    def test_les_parts_de_charge_somment_a_cent_pour_cent(self):
        lignes = self._lignes("2. Sinistralité")
        entetes = [str(c) if c else "" for c in lignes[1]]
        self.assertIn("Part de la charge", entetes)
        colonne = entetes.index("Part de la charge")

        parts = [_nombre(l[colonne]) for l in lignes[2:]]
        parts = [p for p in parts if p is not None]
        self.assertGreater(len(parts), 2)
        self.assertAlmostEqual(
            1.0, sum(parts), delta=0.01,
            msg="Les parts de charge somment a %.4f et non a 1." % sum(parts))


class TestLeWordCombineSOuvreEtPorteDuTexte(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.rapport = TestLeClasseurCombineSOuvreEtPorteSesMontants.rapport

    def _texte(self):
        """⚠️ `doc.paragraphs` IGNORE LES TABLEAUX — piège déjà payé une fois,
        il avait fait conclure à un Word « nettement plus pauvre »."""
        from docx import Document

        doc = Document(io.BytesIO(self.rapport["word_bytes"]))
        morceaux = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for ligne in table.rows:
                morceaux.extend(c.text for c in ligne.cells)
        return "\n".join(morceaux), doc

    def test_le_docx_s_ouvre(self):
        texte, doc = self._texte()
        self.assertGreater(len(doc.paragraphs), 5)
        self.assertGreater(
            len(texte), 1_000,
            "Le Word signe ne porte que %d caracteres de texte, tableaux "
            "compris." % len(texte))

    def test_le_docx_porte_des_tableaux_remplis(self):
        _, doc = self._texte()
        self.assertGreater(len(doc.tables), 2,
                           "Un rapport actuariel sans tableau n'en est pas un.")
        cellules_pleines = sum(
            1 for t in doc.tables for l in t.rows for c in l.cells
            if c.text.strip())
        self.assertGreater(
            cellules_pleines, 20,
            "Seulement %d cellule(s) remplie(s) dans les tableaux du Word."
            % cellules_pleines)


if __name__ == "__main__":
    unittest.main()
